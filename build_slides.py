# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xE8, 0x7A, 0x2B)
GRAY = RGBColor(0x44, 0x44, 0x44)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def box(s, l, t, w, h, fill):
    sp = s.shapes.add_shape(1, l, t, w, h)
    sp.shadow.inherit = False
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    return sp


def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.1):
    tf = s.shapes.add_textbox(l, t, w, h).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4); p.space_before = Pt(0); p.line_spacing = ls
        for txt, size, color, bold in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = "Noto Sans JP"


# ---- Slide 1 ----
s = prs.slides.add_slide(blank)
box(s, 0, 0, SW, Inches(1.2), NAVY)
box(s, 0, Inches(1.2), SW, Inches(0.07), ACCENT)
text(s, Inches(0.6), Inches(0.16), Inches(12.1), Inches(0.9),
     [[("KITAICHI ｜ 地域おこし協力隊のためのクラウドファンディング活用", 23, WHITE, True)],
      [("～ 地域での新しい挑戦を、応援者とともに ～", 14, RGBColor(0xCF,0xDB,0xE8), False)]],
     anchor=MSO_ANCHOR.MIDDLE)

text(s, Inches(0.6), Inches(1.5), Inches(7.4), Inches(0.5), [[("1. KITAICHIの想い", 18, NAVY, True)]])
box(s, Inches(0.6), Inches(2.0), Inches(7.0), Inches(4.95), LIGHT)
text(s, Inches(0.85), Inches(2.2), Inches(6.5), Inches(4.6),
     [[("クラウドファンディングは、単なる資金調達ではなく、地域で新しい挑戦を始める人が、地域内外の応援者と出会い、活動を継続していくための仕組みです。", 13, GRAY, False)],
      [("地域おこし協力隊の皆さまは、任期中に地域の課題や資源に深く関わり、任期後の起業・事業化・定住に向けて挑戦を始める可能性を持った存在です。", 13, GRAY, False)],
      [("一方で、いざ活動を立ち上げると、資金面だけでなく、発信・応援者づくり・見せ方・初期の実績づくりなど、さまざまな壁に直面します。", 13, GRAY, False)],
      [("KITAICHIは、特定サービスの利用を前提とせず、協力隊の皆さまが「地域で挑戦を始める選択肢」としてクラウドファンディングを知り、活かせるかを考えるきっかけを提供します。", 13, GRAY, False)],
      [("自治体の皆さまにとっても、任期後の定着・起業支援・関係人口づくりにつながる情報提供としてご活用いただけます。", 13, NAVY, True)]],
     ls=1.25)

text(s, Inches(7.9), Inches(1.5), Inches(5.0), Inches(0.5), [[("2. クラウドファンディングのメリット", 17, NAVY, True)]])
box(s, Inches(7.9), Inches(2.0), Inches(4.85), Inches(4.95), WHITE)
box(s, Inches(7.9), Inches(2.0), Inches(0.12), Inches(4.95), ACCENT)
merits = [
    ("① 活動を知ってもらうきっかけ", "想いをページに整理し、これまで接点のなかった人にも発信できる。"),
    ("② 応援者・ファンづくり", "共感して応援してくれる人と出会い、任期後を支える基盤になる。"),
    ("③ ニーズの確認", "支援の集まり方で反応を確認。事業化前のテストマーケに。"),
    ("④ 地域内の協力者を増やす", "事業者・自治体・住民との連携が生まれる。"),
    ("⑤ 任期後に向けた初期実績", "資金・支援者数・メディア掲載・協力体制などが実績に。"),
]
runs = []
for t, d in merits:
    runs.append([(t, 12.5, NAVY, True)])
    runs.append([(d, 10.5, GRAY, False)])
text(s, Inches(8.2), Inches(2.2), Inches(4.4), Inches(4.6), runs, ls=1.15)

# ---- Slide 2 ----
s2 = prs.slides.add_slide(blank)
box(s2, 0, 0, SW, Inches(1.0), NAVY)
box(s2, 0, Inches(1.0), SW, Inches(0.07), ACCENT)
text(s2, Inches(0.6), Inches(0.1), Inches(12), Inches(0.85),
     [[("事例紹介とKITAICHIのサポート", 23, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)

text(s2, Inches(0.6), Inches(1.25), Inches(6), Inches(0.45), [[("3. 事例紹介", 18, NAVY, True)]])
box(s2, Inches(0.6), Inches(1.75), Inches(5.6), Inches(3.9), LIGHT)
box(s2, Inches(0.6), Inches(1.75), Inches(0.12), Inches(3.9), ACCENT)
text(s2, Inches(0.85), Inches(1.95), Inches(5.2), Inches(3.6),
     [[("事例A｜北海道・上川町（れいちぇるさん）", 13, NAVY, True)],
      [("「人生をちょっと好きになれる」カフェ＆フォトスタジオ", 11, GRAY, False)],
      [("達成 4,178,638円 / 361人 / 139%", 12, ACCENT, True),
       ("（目標300万円）", 10, GRAY, False)],
      [("", 7, GRAY, False)],
      [("事例B｜北海道・厚沢部町（協力隊・中村和恵さん）", 13, NAVY, True)],
      [("古民家カフェ「tabele」", 11, GRAY, False)],
      [("達成 1,105,000円 / 107人 / 110%", 12, ACCENT, True),
       ("（目標100万円）", 10, GRAY, False)],
      [("", 7, GRAY, False)],
      [("いずれも地域おこし協力隊による北海道の「食・地域づくり」事例。", 10, RGBColor(0x88,0x88,0x88), False)]],
     ls=1.2)

text(s2, Inches(6.7), Inches(1.25), Inches(6), Inches(0.45), [[("4. KITAICHIのサポート内容", 18, NAVY, True)]])
box(s2, Inches(6.7), Inches(1.75), Inches(6.05), Inches(3.9), WHITE)
box(s2, Inches(6.7), Inches(1.75), Inches(0.12), Inches(3.9), ACCENT)
support = [
    ("① 企画整理", "やりたいこと・地域での意味・魅力を整理し成立しやすい企画に。"),
    ("② 目標金額・リターン設計", "金額の考え方、リターン内容・価格、支援導線を一緒に設計。"),
    ("③ ページ構成・文章作成", "構成・タイトル・本文・写真など想いが伝わるページづくり。"),
    ("④ 広報・発信計画", "公開前～ラストスパートまで時期ごとの発信を整理。"),
    ("⑤ 公開後の伴走", "支援状況を見ながら追加施策・ネクストゴール・報告を支援。"),
]
runs = []
for t, d in support:
    runs.append([(t, 12.5, NAVY, True)])
    runs.append([(d, 10.5, GRAY, False)])
text(s2, Inches(7.0), Inches(1.95), Inches(5.6), Inches(3.6), runs, ls=1.15)

box(s2, 0, Inches(6.55), SW, Inches(0.95), NAVY)
text(s2, Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.85),
     [[("研修・定例会・相談会・勉強会などの場で、KITAICHIのクラウドファンディング活用セミナーをぜひご紹介ください。", 13, WHITE, True)]],
     anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

prs.save("/home/user/kitaichi/kitaichi_slides.pptx")
print("saved")
