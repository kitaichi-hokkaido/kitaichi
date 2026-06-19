# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# Colors
NAVY = RGBColor(0x1F, 0x35, 0x5E)
BLUE = RGBColor(0x2E, 0x6D, 0xB4)
LIGHTBLUE = RGBColor(0x4A, 0x90, 0xC2)
RED = RGBColor(0xE2, 0x4A, 0x4A)
ORANGE = RGBColor(0xF2, 0x9A, 0x1B)
GREEN = RGBColor(0x4C, 0xA3, 0x4C)
YELLOW = RGBColor(0xFF, 0xD2, 0x3E)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHTGRAY = RGBColor(0xF0, 0xF1, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x33, 0x33, 0x33)

prs = Presentation()
# A4 portrait
prs.slide_width = Cm(21.0)
prs.slide_height = Cm(29.7)
slide = prs.slides.add_slide(prs.slide_layouts[6])

JPFONT = "Meiryo"

def set_font(run, size, color, bold=False, font=JPFONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn('a:ea'), {'typeface': font})
    rPr.append(ea)

def box(x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=None):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp

def text(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True, space=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    # runs: list of paragraphs, each paragraph is list of (txt,size,color,bold)
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = space
        for (txt, size, color, bold) in para:
            r = p.add_run(); r.text = txt
            set_font(r, size, color, bold)
    return tb

# ===== Background =====
box(0, 0, prs.slide_width, prs.slide_height, fill=WHITE)

# ===== Header tag =====
box(Cm(0.8), Cm(0.6), Cm(9.0), Cm(0.95), fill=NAVY)
text(Cm(0.8), Cm(0.6), Cm(9.0), Cm(0.95),
     [[("挑戦するあなたの、一歩を応援します！", 13, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ===== Main title =====
text(Cm(0.8), Cm(1.7), Cm(11.5), Cm(5.0),
     [[("資金調達の", 30, NAVY, True)],
      [("新しい選択肢、", 30, NAVY, True)],
      [("クラウドファンディング", 30, RED, True)],
      [("というカタチ。", 30, NAVY, True)]],
     space=1.0)

# Hero area (right) - light blue rounded
box(Cm(12.6), Cm(1.3), Cm(7.6), Cm(5.4), fill=RGBColor(0xDB,0xEC,0xF7),
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
# speech bubble
box(Cm(15.6), Cm(1.4), Cm(4.0), Cm(1.6), fill=YELLOW,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(Cm(15.6), Cm(1.4), Cm(4.0), Cm(1.6),
     [[("はじめてでも、", 9, NAVY, True)],
      [("一人じゃないから心強い！", 9, NAVY, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(Cm(12.6), Cm(3.3), Cm(7.6), Cm(3.0),
     [[("👩  🧑  👧  👨", 30, NAVY, False)],
      [("（イラスト差し替え用）", 8, GRAY, False)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Lead text
text(Cm(0.8), Cm(6.7), Cm(11.5), Cm(2.4),
     [[("クラウドファンディングは、資金を集めるだけではありません。", 11, DARK, False)],
      [("あなたの挑戦を多くの人に知ってもらい、", 11, DARK, False)],
      [("共感や応援", 11, RED, True),("を力に変えることができます。", 11, DARK, False)]],
     space=1.15)

# ===== Section 1: 3つの価値 =====
def section_header(y, label):
    w = Cm(11.0); x = (prs.slide_width - w)//2
    h = Cm(0.85)
    # chevron-ish bar
    box(x, y, w, h, fill=NAVY, shape=MSO_SHAPE.PENTAGON)
    box(x, y, w, h, fill=None)  # spacer noop
    text(x, y, w, h, [[(label, 13, WHITE, True)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

sec1_y = Cm(9.3)
section_header(sec1_y, "クラウドファンディングで得られる３つの価値")
# panel
box(Cm(0.8), Cm(10.35), Cm(19.4), Cm(4.5), fill=LIGHTGRAY,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)

values = [
    ("💰", GREEN, "資金調達", "新しい挑戦に必要な資金を集めることができます。"),
    ("📣", RED, "情報発信・認知拡大", "プロジェクトや想いを広く伝え、地域や全国に知ってもらえます。"),
    ("❤️", ORANGE, "ファンづくり", "共感してくれる応援者がファンになり、未来のお客様につながります。"),
]
colw = Cm(6.2)
for i,(icon,c,title,desc) in enumerate(values):
    cx = Cm(1.2) + i*Cm(6.33)
    circ = box(cx, Cm(10.8), Cm(1.5), Cm(1.5), fill=c, shape=MSO_SHAPE.OVAL)
    text(cx, Cm(10.8), Cm(1.5), Cm(1.5), [[(icon, 16, WHITE, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(cx+Cm(1.7), Cm(10.85), colw-Cm(1.7), Cm(1.5),
         [[(title, 12.5, c, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(cx, Cm(12.5), colw, Cm(2.2),
         [[(desc, 10, DARK, False)]], space=1.1)

# ===== Section 2: こんな挑戦に =====
sec2_y = Cm(15.3)
section_header(sec2_y, "こんな挑戦に活用されています")
box(Cm(0.8), Cm(16.35), Cm(19.4), Cm(4.0), fill=LIGHTGRAY,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
uses = [
    ("🏪","開業・起業","夢だったお店をオープンしたい！"),
    ("🧴","商品・サービス開発","こだわりの商品をカタチにしたい！"),
    ("📈","事業拡大","もっと多くの人に届けたい！"),
    ("💡","設備・改装","居心地の良いお店にしたい！"),
    ("🏞️","地域・観光・まちづくり","地域を元気にする活動を広げたい！"),
    ("🔍","テストマーケティング","新しいアイデアの反応を確かめたい！"),
]
cw = Cm(3.1)
for i,(icon,title,desc) in enumerate(uses):
    cx = Cm(1.1) + i*Cm(3.18)
    box(cx, Cm(16.6), cw, Cm(3.5), fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(cx, Cm(16.75), cw, Cm(0.6), [[(title, 8.5, NAVY, True)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(cx, Cm(17.4), cw, Cm(1.3), [[(icon, 22, NAVY, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(cx, Cm(18.7), cw, Cm(1.3), [[(desc, 8, DARK, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=1.05)

# ===== Section 3: お悩み / サポート =====
sec3_y = Cm(20.8)
# left panel - worries
box(Cm(0.8), sec3_y, Cm(9.3), Cm(4.5), fill=LIGHTGRAY,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(Cm(0.8), sec3_y+Cm(0.2), Cm(9.3), Cm(0.7),
     [[("こんなお悩みありませんか？", 12, NAVY, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
worries = ["クラウドファンディングって難しそう…",
           "一人で準備するのは負担が大きい…",
           "どうやって支援を集めたらいいの？",
           "ページ作りや発信に自信がない…",
           "失敗したらどうしよう…"]
text(Cm(2.2), sec3_y+Cm(1.0), Cm(7.8), Cm(3.4),
     [[("☑  "+w, 9.5, DARK, False)] for w in worries], space=1.25)

# right panel - support
box(Cm(10.4), sec3_y, Cm(9.8), Cm(4.5), fill=RGBColor(0xDB,0xEC,0xF7),
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(Cm(10.4), sec3_y+Cm(0.2), Cm(9.8), Cm(0.7),
     [[("KITAICHIが伴走サポートします！", 12, NAVY, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
sup = ["企画・設計から公開後まで、しっかり伴走",
       "魅力が伝わるページ制作や発信をサポート",
       "クラウドファンディングが初めての方も安心"]
text(Cm(10.9), sec3_y+Cm(1.0), Cm(8.8), Cm(2.4),
     [[("✓  "+s, 9.5, NAVY, True)] for s in sup], space=1.3)
text(Cm(10.9), sec3_y+Cm(3.4), Cm(8.8), Cm(0.9),
     [[("あなたの想いを、カタチにするお手伝いをします。", 9.5, DARK, False)]])

# arrow between
box(Cm(9.9), sec3_y+Cm(1.6), Cm(0.9), Cm(1.2), fill=BLUE, shape=MSO_SHAPE.CHEVRON)

# ===== CTA band =====
cta_y = Cm(25.6)
box(Cm(0.4), cta_y, Cm(20.2), Cm(2.5), fill=YELLOW,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(Cm(0.8), cta_y+Cm(0.55), Cm(1.8), Cm(1.4), fill=NAVY,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(Cm(0.8), cta_y+Cm(0.55), Cm(1.8), Cm(1.4), [[("まずは", 11, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(Cm(2.8), cta_y+Cm(0.3), Cm(11.0), Cm(2.0),
     [[("無料で相談してみませんか？", 17, NAVY, True)],
      [("「何から始めればいいかわからない…」という方も大歓迎！", 9, DARK, False)],
      [("あなたの挑戦に合った方法をご提案します。", 9, DARK, False)]],
     space=1.1)
box(Cm(16.3), cta_y+Cm(0.45), Cm(3.9), Cm(1.6), fill=NAVY,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(Cm(16.3), cta_y+Cm(0.45), Cm(3.9), Cm(1.6),
     [[("お気軽に", 12, WHITE, True)],[("ご連絡ください！", 12, WHITE, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ===== Footer =====
foot_y = Cm(28.3)
text(Cm(0.8), foot_y, Cm(7.0), Cm(1.3),
     [[("あなたの挑戦を、応援でつなぐパートナー", 8, GRAY, False)],
      [("株式会社 KITAICHI", 14, NAVY, True)]], space=1.1)
text(Cm(8.0), foot_y, Cm(6.5), Cm(1.3),
     [[("📞 03-6804-2813", 9, DARK, False)],
      [("✉ info@kitaichi.co.jp", 9, DARK, False)],
      [("🌐 https://kitaichi.co.jp/", 9, DARK, False)]], space=1.05)
text(Cm(14.6), foot_y, Cm(5.0), Cm(1.3),
     [[("企画からページ制作、リターン設計、", 7.5, GRAY, False)],
      [("プロモーションまでワンストップでサポート！", 7.5, GRAY, False)],
      [("あなたの挑戦に伴走します。", 7.5, GRAY, False)]], space=1.05)

prs.save("KITAICHI_crowdfunding.pptx")
print("saved")
