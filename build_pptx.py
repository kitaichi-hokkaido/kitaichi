# -*- coding: utf-8 -*-
"""ふるさと住民登録制度 企画運営サポート提案 .pptx generator"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# palette
NAVY   = RGBColor(0x13,0x31,0x4F)
NAVY2  = RGBColor(0x1F,0x4D,0x77)
BLUE   = RGBColor(0x2F,0x80,0xC4)
SKY    = RGBColor(0x5A,0xA9,0xE6)
ACCENT = RGBColor(0xF2,0x76,0x2E)
ACCENT2= RGBColor(0xFF,0xB0,0x66)
INK    = RGBColor(0x1B,0x27,0x33)
SUB    = RGBColor(0x5D,0x6F,0x80)
LIGHT  = RGBColor(0xF2,0xF6,0xFA)
LINE   = RGBColor(0xE2,0xEA,0xF1)
WHITE  = RGBColor(0xFF,0xFF,0xFF)
GREEN  = RGBColor(0x1F,0x8A,0x4D)
GREY   = RGBColor(0xCD,0xD9,0xE3)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
FONT = "Meiryo"

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def _set_font(run, size, color, bold, font=FONT):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.name = font
    # east asian font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', font)

def textbox(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """lines: list of (text, size, color, bold) or list of such => paragraphs."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first = True
    for para in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if isinstance(para, tuple):
            para = [para]
        for seg in para:
            text, size, color, bold = seg[0], seg[1], seg[2], seg[3]
            r = p.add_run(); r.text = text
            _set_font(r, size, color, bold)
        # line spacing
        p.line_spacing = 1.15
    return tb

def rect(s, x, y, w, h, fill, line=None, rounded=False, line_w=1.0):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        try: shp.adjustments[0] = 0.08
        except Exception: pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp

def shape_text(shp, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left=Inches(0.08); tf.margin_right=Inches(0.08)
    tf.margin_top=Inches(0.04); tf.margin_bottom=Inches(0.04)
    first=True
    for para in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first=False; p.alignment=align; p.line_spacing=1.1
        if isinstance(para, tuple): para=[para]
        for seg in para:
            r=p.add_run(); r.text=seg[0]; _set_font(r, seg[1], seg[2], seg[3])

def oval(s,x,y,w,h,fill):
    o=s.shapes.add_shape(MSO_SHAPE.OVAL,x,y,w,h)
    o.fill.solid(); o.fill.fore_color.rgb=fill; o.line.fill.background()
    o.shadow.inherit=False; return o

def chevron(s,x,y,w,h,fill):
    c=s.shapes.add_shape(MSO_SHAPE.CHEVRON,x,y,w,h)
    c.fill.solid(); c.fill.fore_color.rgb=fill; c.line.fill.background()
    c.shadow.inherit=False; return c

def footer(s, n, dark=False):
    col = GREY if dark else RGBColor(0xAA,0xB8,0xC4)
    textbox(s, Inches(0.5), Inches(7.05), Inches(3), Inches(0.3),
            [("KITAICHI", 10, col, True)])
    textbox(s, Inches(11.8), Inches(7.05), Inches(1.0), Inches(0.3),
            [("%02d"%n, 10, col, False)], align=PP_ALIGN.RIGHT)

def kicker(s, text, x=Inches(0.83), y=Inches(0.62), color=ACCENT):
    # small bar + label
    rect(s, x, y+Inches(0.08), Inches(0.32), Inches(0.05), color)
    textbox(s, x+Inches(0.42), y-Inches(0.02), Inches(8), Inches(0.35),
            [(text, 13, color, True)])

LX = Inches(0.83)  # left margin

# ============ 1 COVER ============
s = slide(); bg(s, NAVY)
# decorative circles
c1=oval(s, Inches(9.3), Inches(-2.4), Inches(5.5), Inches(5.5), NAVY2)
c2=oval(s, Inches(-1.6), Inches(4.3), Inches(4.2), Inches(4.2), NAVY2)
textbox(s, LX, Inches(1.55), Inches(11), Inches(0.4),
        [("ふるさと住民登録制度　企画運営サポートのご提案", 15, ACCENT2, True)])
textbox(s, LX, Inches(2.25), Inches(11.5), Inches(2.2),
        [[("地域の応援者を増やし、", 42, WHITE, True)],
         [("継続的な「関わり」をつくる。", 42, WHITE, True)]])
textbox(s, LX, Inches(4.5), Inches(11), Inches(1.0),
        [[("～ クラウドファンディングで培った「地域と人をつなぐ力」を、", 17, RGBColor(0xBC,0xD2,0xE8), False)],
         [("ふるさと住民登録制度の運用へ ～", 17, RGBColor(0xBC,0xD2,0xE8), False)]])
textbox(s, Inches(8.8), Inches(6.2), Inches(4.0), Inches(0.9),
        [[("KITAICHI", 22, WHITE, True)],[("株式会社KITAICHI", 12, RGBColor(0x9F,0xB6,0xCD), False)]],
        align=PP_ALIGN.RIGHT)

# ============ 2 SUMMARY ============
s=slide(); bg(s,WHITE)
kicker(s,"このご提案を一言でいうと")
textbox(s, LX, Inches(1.15), Inches(11.5), Inches(1.3),
        [[("制度の「企画 → 募集 → 担い手活動」まで、", 28, NAVY, True)],
         [("KITAICHIがまるごと伴走します。", 28, NAVY, True)]])
labels=["① 企画設計","② 募集支援","③ 担い手活動"]
bx=Inches(1.2); bw=Inches(2.9); gap=Inches(0.55); by=Inches(3.0)
for i,t in enumerate(labels):
    x=Emu(int(bx)+i*(int(bw)+int(gap)))
    card=rect(s,x,by,bw,Inches(1.5),ACCENT,rounded=True)
    shape_text(card,[(t,20,WHITE,True)])
    if i<2:
        ar=chevron(s,Emu(int(x)+int(bw)+Inches(0.08)),by+Inches(0.5),Inches(0.42),Inches(0.5),ACCENT2)
co=rect(s,LX,Inches(5.3),Inches(11.6),Inches(1.25),RGBColor(0xFF,0xF6,0xEE),line=RGBColor(0xF6,0xD3,0xB3),rounded=True,line_w=1.5)
shape_text(co,[[("費用は特別交付税措置の範囲内で検討可能です（総務省確認済み）。",18,RGBColor(0xB4,0x57,0x1D),True)],
              [("まずは雑談レベルからでも大丈夫です。",16,RGBColor(0xB4,0x57,0x1D),False)]],align=PP_ALIGN.LEFT)
footer(s,2)

# ============ 3 制度の本質 (dark, compare) ============
s=slide(); bg(s,NAVY)
kicker(s,"ふるさと住民登録制度とは",color=ACCENT2)
textbox(s, LX, Inches(1.15), Inches(11.5), Inches(1.4),
        [[("「登録」ではなく、", 30, WHITE, True)],
         [("地域との「関わり」を育てる仕組み。", 30, WHITE, True)]])
textbox(s, LX, Inches(2.95), Inches(11.5), Inches(0.8),
        [("地域を応援したい人・関わりたい人・未来に参加したい人とつながり、継続的な関係を育てる制度です。",16,RGBColor(0xCD,0xDD,0xF0),False)])
# two panes
p1=rect(s,Inches(1.2),Inches(4.2),Inches(4.7),Inches(2.0),NAVY2,rounded=True)
shape_text(p1,[[("ふるさと納税",18,RGBColor(0xCF,0xE0,0xF0),True)],[("",8,WHITE,False)],
              [("「寄附」という形で",16,WHITE,False)],[("地域との接点を広げた",16,WHITE,False)]])
ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(6.15),Inches(4.95),Inches(0.95),Inches(0.5))
ar.fill.solid(); ar.fill.fore_color.rgb=ACCENT; ar.line.fill.background(); ar.shadow.inherit=False
p2=rect(s,Inches(7.35),Inches(4.2),Inches(4.7),Inches(2.0),RGBColor(0x3A,0x2A,0x1E),rounded=True)
shape_text(p2,[[("ふるさと住民登録制度",18,ACCENT2,True)],[("",8,WHITE,False)],
              [("「関わり」という形で",16,RGBColor(0xFF,0xE2,0xCC),False)],[("地域の応援者を増やす",16,RGBColor(0xFF,0xE2,0xCC),False)]])
footer(s,3,dark=True)

# ============ 4 不安 ============
s=slide(); bg(s,WHITE)
kicker(s,"運用検討にあたり")
textbox(s, LX, Inches(1.1), Inches(11), Inches(0.8),[("こんな不安はありませんか？",32,NAVY,True)])
probs=[("不安 01","この制度をどう活用すれば\nよいかわからない"),
       ("不安 02","通常業務で手一杯で、企画や\n運用まで手が回らない"),
       ("不安 03","担い手活動の具体的な\nイメージがわかない")]
px=Inches(1.0); pw=Inches(3.7); pg=Inches(0.35); py=Inches(2.4)
for i,(q,t) in enumerate(probs):
    x=Emu(int(px)+i*(int(pw)+int(pg)))
    card=rect(s,x,py,pw,Inches(2.6),RGBColor(0xF6,0xF9,0xFC),line=LINE,rounded=True)
    ic=rect(s,Emu(int(x)+int(Inches(0.35))),py+Inches(0.35),Inches(0.75),Inches(0.75),RGBColor(0xE7,0xEE,0xF5),rounded=True)
    shape_text(ic,[("?",30,NAVY2,True)])
    textbox(s,Emu(int(x)+int(Inches(0.35))),py+Inches(1.3),Inches(3.0),Inches(0.3),[(q,12,SUB,True)])
    tb=textbox(s,Emu(int(x)+int(Inches(0.35))),py+Inches(1.65),Inches(3.0),Inches(0.9),
               [(seg,17,INK,False) for seg in t.split("\n")])
footer(s,4)

# ============ 5 伴走 (dark center) ============
s=slide(); bg(s,NAVY)
oval(s,Inches(-1.5),Inches(-1.5),Inches(4),Inches(4),NAVY2)
textbox(s,Inches(0),Inches(2.3),SW,Inches(0.4),[("その不安に",14,ACCENT2,True)],align=PP_ALIGN.CENTER)
textbox(s,Inches(0),Inches(2.85),SW,Inches(1.0),[("KITAICHIが、伴走します。",46,WHITE,True)],align=PP_ALIGN.CENTER)
textbox(s,Inches(2.4),Inches(4.3),Inches(8.5),Inches(1.2),
        [("ふるさと住民登録制度は、クラウドファンディングに非常に近い構造です。制度設計・募集・担い手活動の企画運営まで、まるごとお任せください。",17,RGBColor(0xCD,0xDD,0xF0),False)],align=PP_ALIGN.CENTER)
footer(s,5,dark=True)

# ============ 6 CFと似ている (compare) ============
s=slide(); bg(s,WHITE)
kicker(s,"なぜ伴走できるのか")
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(1.3),
        [[("ふるさと住民登録制度は、",26,NAVY,True)],[("クラウドファンディングとよく似ています。",26,NAVY,True)]])
cf=["取り組みの背景や想いを伝える","地域の魅力や課題を届ける","共感してくれる人を増やす","応援や参加を募る","継続的な関係をつくる"]
fr=["地域の想いやストーリーを届ける","地域の魅力をわかりやすく伝える","共感する登録者を増やす","担い手活動への参加を募る","登録後も継続的な関係を育てる"]
p1=rect(s,Inches(0.9),Inches(2.75),Inches(5.0),Inches(3.4),RGBColor(0xEA,0xF3,0xFB),rounded=True)
textbox(s,Inches(1.2),Inches(3.0),Inches(4.5),Inches(0.4),[("クラウドファンディング",18,NAVY2,True)])
textbox(s,Inches(1.2),Inches(3.55),Inches(4.5),Inches(2.5),
        [[("✓ ",15,BLUE,True),(t,15,INK,False)] for t in cf])
p2=rect(s,Inches(7.45),Inches(2.75),Inches(5.0),Inches(3.4),RGBColor(0xFF,0xF1,0xE6),rounded=True)
textbox(s,Inches(7.75),Inches(3.0),Inches(4.5),Inches(0.4),[("ふるさと住民登録制度",18,RGBColor(0xC2,0x61,0x1D),True)])
textbox(s,Inches(7.75),Inches(3.55),Inches(4.5),Inches(2.5),
        [[("✓ ",15,ACCENT,True),(t,15,INK,False)] for t in fr])
eq=oval(s,Inches(6.1),Inches(4.0),Inches(1.0),Inches(1.0),NAVY); shape_text(eq,[("≒",30,WHITE,True)])
textbox(s,Inches(0),Inches(6.4),SW,Inches(0.4),
        [[("＝ CFのノウハウが、そのまま担い手活動に活きます。",17,NAVY,True)]],align=PP_ALIGN.CENTER)
footer(s,6)

# ============ 7 section ============
def section(n, label, title):
    s=slide(); bg(s,NAVY)
    oval(s,Inches(9.5),Inches(3.5),Inches(5),Inches(5),NAVY2)
    kicker(s,label,y=Inches(2.8),color=ACCENT2)
    textbox(s,LX,Inches(3.25),Inches(11.5),Inches(1.2),[(title,44,WHITE,True)])
    footer(s,n,dark=True)
section(7,"SECTION 01","なぜKITAICHIが支援できるのか")

# ============ 8 実績 stats ============
s=slide(); bg(s,WHITE)
kicker(s,"私たちの実績")
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(0.8),[("地域の挑戦を、全国へ届けてきました。",30,NAVY,True)])
textbox(s,LX,Inches(2.0),Inches(11.5),Inches(0.9),
        [("単なる資金調達ではなく、挑戦する人の想いを整理し、ストーリーとして届け、共感する応援者とつなげること。それが私たちの仕事です。",16,SUB,False)])
stats=[("35","件","サポート実績"),("6,488","万円","累計支援総額"),("4,060","人","累計支援者数")]
sx=Inches(1.0); sw_=Inches(3.7); sg=Inches(0.35); sy=Inches(3.2)
for i,(big,unit,cap) in enumerate(stats):
    x=Emu(int(sx)+i*(int(sw_)+int(sg)))
    card=rect(s,x,sy,sw_,Inches(2.1),RGBColor(0xF1,0xF6,0xFB),line=LINE,rounded=True)
    rect(s,x,sy,sw_,Inches(0.12),ACCENT)
    textbox(s,x,sy+Inches(0.5),sw_,Inches(0.9),[[(big,44,NAVY,True),(unit,20,NAVY,True)]],align=PP_ALIGN.CENTER)
    textbox(s,x,sy+Inches(1.5),sw_,Inches(0.4),[(cap,15,SUB,False)],align=PP_ALIGN.CENTER)
textbox(s,LX,Inches(5.7),Inches(11.5),Inches(0.5),
        [("※ KITAICHIはCAMPFIREのパートナー企業として、まちづくりプロジェクトを専門にサポートしています。",14,SUB,False)])
footer(s,8)

# ============ 9 スケール (dark stats) ============
s=slide(); bg(s,NAVY)
kicker(s,"私たちの強み",color=ACCENT2)
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(1.3),
        [[("地域の取り組みを、",30,WHITE,True)],
         [("「全国」へ届ける発信網があります。",30,ACCENT2,True)]])
sc=[("850","件","令和7年度 GCF業務委託で\n立ち上げ支援したプロジェクト"),
    ("全国","","CAMPFIREパートナーとして\n地域の挑戦を全国の支援者へ"),
    ("4,060","人","「応援したい人」と\nつながってきた実績")]
sx=Inches(1.0); sw_=Inches(3.7); sg=Inches(0.35); sy=Inches(3.1)
for i,(big,unit,cap) in enumerate(sc):
    x=Emu(int(sx)+i*(int(sw_)+int(sg)))
    card=rect(s,x,sy,sw_,Inches(2.0),NAVY2,rounded=True)
    textbox(s,x,sy+Inches(0.3),sw_,Inches(0.9),[[(big,40,WHITE,True),(unit,18,WHITE,True)]],align=PP_ALIGN.CENTER)
    textbox(s,x,sy+Inches(1.15),sw_,Inches(0.8),[(seg,13,RGBColor(0xBC,0xD2,0xE8),False) for seg in cap.split("\n")],align=PP_ALIGN.CENTER)
co=rect(s,LX,Inches(5.5),Inches(11.6),Inches(1.05),RGBColor(0x3A,0x2A,0x1E),rounded=True)
shape_text(co,[[("ふるさと納税が「寄附」で全国に広がったように、本制度は「関わり」で全国に広がる可能性があります。",16,RGBColor(0xFF,0xD9,0xBB),True)]],align=PP_ALIGN.LEFT)
footer(s,9,dark=True)

# ============ 10 イベント実績 ============
s=slide(); bg(s,WHITE)
kicker(s,"実績")
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(0.8),[("セミナー・地域イベントの企画運営も。",30,NAVY,True)])
evs=["夕張市農協の\n婚活イベントの企画運営","大学生を集めた\nリアルセミナーの開催","地域創生イベントへの\n登壇・開催"]
ex=Inches(1.0); ew=Inches(3.7); eg=Inches(0.35); ey=Inches(2.4)
for i,t in enumerate(evs):
    x=Emu(int(ex)+i*(int(ew)+int(eg)))
    card=rect(s,x,ey,ew,Inches(2.0),RGBColor(0xF6,0xF9,0xFC),line=LINE,rounded=True)
    ic=rect(s,Emu(int(x)+int(Inches(0.35))),ey+Inches(0.3),Inches(0.7),Inches(0.7),RGBColor(0xEA,0xF3,0xFB),rounded=True)
    shape_text(ic,[("●",20,BLUE,True)])
    textbox(s,Emu(int(x)+int(Inches(0.35))),ey+Inches(1.15),Inches(3.0),Inches(0.8),
            [(seg,16,INK,True) for seg in t.split("\n")])
textbox(s,LX,Inches(4.9),Inches(11.5),Inches(0.8),
        [("情報発信から、リアルな場づくり・ファンミーティングの開催まで。「人を集め、つなげる」現場の経験が豊富です。",17,SUB,False)])
footer(s,10)

# ============ 11 section ============
section(11,"SECTION 02","ご提案内容")

# ============ 12 ジャーニー ============
s=slide(); bg(s,WHITE)
kicker(s,"私たちが目指すこと")
textbox(s,LX,Inches(1.1),Inches(11),Inches(0.7),[("「登録」で終わらせない。",30,NAVY,True)])
textbox(s,LX,Inches(1.95),Inches(11.5),Inches(0.5),
        [("登録者数を増やすだけでなく、登録した人が地域に関わり続ける流れをつくります。",16,SUB,False)])
steps=[("登録",GREY,RGBColor(0x39,0x50,0x5F)),("共感",SKY,WHITE),("参加",BLUE,WHITE),
       ("交流",NAVY2,WHITE),("継続的な関係",NAVY,WHITE),("担い手・\n応援者へ",ACCENT,WHITE)]
jx=Inches(0.7); jw=Inches(1.75); jg=Inches(0.28); jy=Inches(3.0)
for i,(t,fill,tc) in enumerate(steps):
    x=Emu(int(jx)+i*(int(jw)+int(jg)))
    card=rect(s,x,jy,jw,Inches(1.3),fill,rounded=True)
    shape_text(card,[(seg,16,tc,True) for seg in t.split("\n")])
    if i<len(steps)-1:
        textbox(s,Emu(int(x)+int(jw)+int(Inches(0.01))),jy+Inches(0.45),Inches(0.26),Inches(0.4),
                [("→",18,ACCENT,True)],align=PP_ALIGN.CENTER)
co=rect(s,LX,Inches(5.0),Inches(11.6),Inches(1.1),RGBColor(0xFF,0xF6,0xEE),line=RGBColor(0xF6,0xD3,0xB3),rounded=True,line_w=1.5)
shape_text(co,[[("この流れを設計することで、制度を「地域の未来につながる仕組み」へ育てていきます。",17,RGBColor(0xB4,0x57,0x1D),True)]],align=PP_ALIGN.LEFT)
footer(s,12)

# ============ 13 3ステップ ============
s=slide(); bg(s,WHITE)
kicker(s,"サポートの全体像")
textbox(s,LX,Inches(1.1),Inches(11),Inches(0.7),[("3つのステップで伴走します。",30,NAVY,True)])
sup=[("①","制度活用の企画設計","地域課題や強みを整理し、「どう活用するか」を一緒に検討します。"),
     ("②","登録者の募集支援","制度を多くの人に知ってもらい、登録につなげる情報発信を支援します。"),
     ("③","担い手活動の運営支援","地域に合わせて担い手活動を企画し、運営まで伴走します。")]
sx=Inches(1.0); sw_=Inches(3.7); sg=Inches(0.35); sy=Inches(2.4)
for i,(no,ti,de) in enumerate(sup):
    x=Emu(int(sx)+i*(int(sw_)+int(sg)))
    card=rect(s,x,sy,sw_,Inches(3.0),RGBColor(0xF1,0xF6,0xFB),line=LINE,rounded=True)
    ic=oval(s,Emu(int(x)+int(Inches(0.4))),sy+Inches(0.35),Inches(0.8),Inches(0.8),ACCENT)
    shape_text(ic,[(no,24,WHITE,True)])
    textbox(s,Emu(int(x)+int(Inches(0.4))),sy+Inches(1.4),Inches(3.0),Inches(0.5),[(ti,19,NAVY,True)])
    textbox(s,Emu(int(x)+int(Inches(0.4))),sy+Inches(2.0),Inches(3.0),Inches(0.9),[(de,14,INK,False)])
footer(s,13)

# ============ 14 役割分担 matrix ============
s=slide(); bg(s,WHITE)
kicker(s,"「手が回らない」をそのままにしません")
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(0.7),[("自治体の負担は、KITAICHIが巻き取ります。",28,NAVY,True)])
rows=[("工程","KITAICHIが担う","自治体にお願いする"),
      ("① 企画","コンセプト設計・たたき台の作成","地域の想いの共有・方向性の決定"),
      ("② 募集","ページ・チラシ・SNS発信の制作","最終確認・地域情報の提供"),
      ("③ 担い手活動","企画・運営・当日対応","参加・地域内の橋渡し")]
tx=Inches(0.9); ty=Inches(2.2)
c0=Inches(2.6); c1=Inches(4.7); c2=Inches(4.7); rh=Inches(0.95); hh=Inches(0.7)
for ri,row in enumerate(rows):
    yy=Emu(int(ty)+ (0 if ri==0 else int(hh)+(ri-1)*int(rh)))
    h = hh if ri==0 else rh
    widths=[c0,c1,c2]; xx=int(tx)
    for ci,cell in enumerate(row):
        w=widths[ci]
        if ri==0:
            cellfill=NAVY; tc=WHITE; bold=True; sz=15
        elif ci==0:
            cellfill=RGBColor(0xF3,0xF7,0xFB); tc=NAVY; bold=True; sz=15
        elif ci==1:
            cellfill=WHITE; tc=INK; bold=False; sz=15
        else:
            cellfill=WHITE; tc=INK; bold=False; sz=15
        cc=rect(s,Emu(xx),yy,w,h,cellfill,line=LINE)
        shape_text(cc,[(cell,sz,tc,bold)],align=PP_ALIGN.CENTER if ri==0 or ci==0 else PP_ALIGN.LEFT)
        xx+=int(w)
textbox(s,Inches(0),Inches(6.3),SW,Inches(0.4),
        [[("自治体の皆さまは「判断」と「地域とのつなぎ役」に集中していただけます。",16,SUB,True)]],align=PP_ALIGN.CENTER)
footer(s,14)

# ============ generic detail slide (15,16,17) ============
def detail(n, tag, title, lead, items, cols=3, note=None):
    s=slide(); bg(s,WHITE)
    kicker(s,tag)
    textbox(s,LX,Inches(1.05),Inches(11.5),Inches(0.6),[(title,28,NAVY,True)])
    textbox(s,LX,Inches(1.85),Inches(11.5),Inches(0.6),[(lead,16,SUB,False)])
    gx=Inches(0.9); gy=Inches(2.7)
    gw=Inches(3.83) if cols==3 else Inches(5.75)
    gg=Inches(0.18)
    rows=(len(items)+cols-1)//cols
    ih=Inches(0.82) if not note else Inches(0.72)
    for idx,it in enumerate(items):
        r=idx//cols; c=idx%cols
        x=Emu(int(gx)+c*(int(gw)+int(gg)))
        y=Emu(int(gy)+r*(int(ih)+int(Inches(0.16))))
        card=rect(s,x,y,gw,ih,RGBColor(0xF6,0xF9,0xFC),line=LINE,rounded=True)
        dot=rect(s,Emu(int(x)+int(Inches(0.2))),Emu(int(y)+int(Inches(0.18))),Inches(0.36),Inches(0.36),RGBColor(0xEA,0xF3,0xFB),rounded=True)
        shape_text(dot,[("●",12,BLUE,True)])
        textbox(s,Emu(int(x)+int(Inches(0.7))),Emu(int(y)+int(Inches(0.12))),Emu(int(gw)-int(Inches(0.85))),ih,
                [(it,14,INK,False)],anchor=MSO_ANCHOR.MIDDLE)
    if note:
        co=rect(s,LX,Inches(5.85),Inches(11.6),Inches(0.85),RGBColor(0xFF,0xF6,0xEE),line=RGBColor(0xF6,0xD3,0xB3),rounded=True,line_w=1.5)
        shape_text(co,[[(note,15,RGBColor(0xB4,0x57,0x1D),True)]],align=PP_ALIGN.LEFT)
    footer(s,n)

detail(15,"サポート内容 ①","制度活用の企画設計",
       "自治体ごとの地域課題や強みを整理し、制度をどう活用するかを一緒に検討します。",
       ["制度活用のコンセプト設計","どんな人に登録してほしいかの整理","地域の魅力や課題の言語化",
        "担い手活動の方向性整理","ベーシック／プレミアム登録者の関係設計","登録者に伝えるメッセージ設計"],cols=3)
detail(16,"サポート内容 ②","登録者の募集支援",
       "制度を多くの人に知ってもらい、登録につなげるための情報発信を支援します。",
       ["募集ページの構成づくり","パンフレット・チラシの制作支援","SNS・メール発信内容の設計",
        "CF的な募集導線の設計","地域外の応援者に届くメッセージ","プロモーション戦略の整理"],cols=3,
       note="「登録してください」ではなく、「この地域の未来に関わってください」というメッセージにすることがポイントです。")
detail(17,"サポート内容 ③","担い手活動の企画・運営支援",
       "自治体ごとに内容は大きく変わります。地域に合わせて、例えばこんなことを一緒に企画します。",
       ["地域ファンミーティング","新商品・特産品のモニター会","商品開発アイデア募集","空き家活用アイデア募集",
        "地域おこし協力隊・高校生との交流","町長・村長・地域事業者との交流会","農業・漁業・林業・酒蔵の体験","まちづくり会議へのオブザーバー参加"],cols=2)

# ============ 18 参加の階段 ============
s=slide(); bg(s,WHITE)
kicker(s,"設計の考え方")
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(0.7),[("担い手活動は「参加の階段」で設計します。",28,NAVY,True)])
levels=[("LEVEL 1","知る","SNS・メールで活動を知る",GREY,RGBColor(0x39,0x50,0x5F),1.0),
        ("LEVEL 2","応援する","発信に協力／投票／商品購入",SKY,WHITE,1.5),
        ("LEVEL 3","交流する","交流会・ファンミーティング",BLUE,WHITE,2.0),
        ("LEVEL 4","参加する","現地イベント／モニター・企画会議",NAVY2,WHITE,2.5),
        ("LEVEL 5","担い手になる","継続参加・共創／移住・起業へ",ACCENT,WHITE,3.0)]
bx=Inches(0.8); bw=Inches(2.3); bg_=Inches(0.18); baseY=Inches(6.0)
for i,(lv,nm,cap,fill,tc,hh) in enumerate(levels):
    x=Emu(int(bx)+i*(int(bw)+int(bg_)))
    h=Inches(hh)
    y=Emu(int(baseY)-int(h))
    bar=rect(s,x,y,bw,h,fill,rounded=True)
    shape_text(bar,[[(lv,11,tc,True)],[(nm,18,tc,True)]],anchor=MSO_ANCHOR.TOP)
    textbox(s,x,Emu(int(baseY)+int(Inches(0.1))),bw,Inches(0.8),[(cap,11,SUB,False)],align=PP_ALIGN.CENTER)
textbox(s,Inches(0),Inches(6.75),SW,Inches(0.35),
        [[("登録者が少しずつ関わりを深められる設計が、継続のカギです。",15,SUB,True)]],align=PP_ALIGN.CENTER)
footer(s,18)

# ============ 19 事例 (dark timeline) ============
s=slide(); bg(s,NAVY)
kicker(s,"イメージ",color=ACCENT2)
textbox(s,LX,Inches(1.1),Inches(11.5),Inches(0.7),[("ある「応援者」が「担い手」に育つまで。",30,WHITE,True)])
story=[("Lv.1","知る","都内在住のAさん。SNSで御地の情報を受け取るだけの存在だった。",NAVY2),
       ("Lv.4","参加する","特産品のモニター会に申し込み、初めて地域とリアルに関わる。",NAVY2),
       ("Lv.4+","貢献する","商品開発アイデア募集で意見が採用。役に立つ実感を得る。",NAVY2),
       ("Lv.5","担い手に","翌年には企画会議の常連に。「いつか移住したい」と話すように。",ACCENT)]
cx=Inches(0.8); cw=Inches(2.85); cg=Inches(0.25); cy=Inches(2.5)
for i,(lv,nm,tx_,col) in enumerate(story):
    x=Emu(int(cx)+i*(int(cw)+int(cg)))
    dot=oval(s,x,cy,Inches(0.8),Inches(0.8),col); shape_text(dot,[(lv,13,WHITE,True)])
    textbox(s,x,cy+Inches(0.95),cw,Inches(0.35),[(nm,13,ACCENT2,True)])
    textbox(s,x,cy+Inches(1.35),cw,Inches(1.8),[(tx_,15,RGBColor(0xCD,0xDD,0xF0),False)])
    if i<3:
        textbox(s,Emu(int(x)+int(cw)),cy+Inches(0.15),Inches(0.25),Inches(0.4),[("→",20,ACCENT,True)],align=PP_ALIGN.CENTER)
textbox(s,LX,Inches(5.8),Inches(11.5),Inches(0.5),
        [[("—— こうして登録者が、少しずつ地域の担い手へ育っていきます。",16,ACCENT2,True)]])
footer(s,19,dark=True)

# ============ 20 効果 ============
s=slide(); bg(s,WHITE)
kicker(s,"期待できる効果")
textbox(s,LX,Inches(1.1),Inches(11),Inches(0.7),[("制度を活用することで、こんな効果が。",30,NAVY,True)])
effs=["地域外の応援者を可視化できる","継続的に届けられる関係人口が増える","イベント・活動への参加者が増える",
      "地域の担い手不足の解消につながる","ふるさと納税・CFとの連携がしやすい","将来的な移住・関係人口拡大へ"]
gx=Inches(1.0); gy=Inches(2.4); gw=Inches(3.7); gg=Inches(0.35); ih=Inches(1.7)
for i,t in enumerate(effs):
    r=i//3; c=i%3
    x=Emu(int(gx)+c*(int(gw)+int(gg)))
    y=Emu(int(gy)+r*(int(ih)+int(Inches(0.3))))
    card=rect(s,x,y,gw,ih,RGBColor(0xF6,0xF9,0xFC),line=LINE,rounded=True)
    ic=rect(s,Emu(int(x)+int(Inches(0.3))),Emu(int(y)+int(Inches(0.3))),Inches(0.6),Inches(0.6),RGBColor(0xEA,0xF3,0xFB),rounded=True)
    shape_text(ic,[("✓",20,BLUE,True)])
    textbox(s,Emu(int(x)+int(Inches(0.3))),Emu(int(y)+int(Inches(1.0))),Emu(int(gw)-int(Inches(0.6))),Inches(0.6),[(t,15,INK,True)])
footer(s,20)

# ============ 21 財源 ============
s=slide(); bg(s,WHITE)
kicker(s,"財源・ご相談について")
textbox(s,LX,Inches(1.1),Inches(11),Inches(0.7),[("活用できる財源があります。",30,NAVY,True)])
fin=[("制度の準備・運用","上限 1,000万円","事業費1,000万円を上限に、50％×財政力補正により特別交付税措置。本業務への活用も可能です。"),
     ("ふるさと住民コーディネーターの設置","500万円","500万円の特別交付税措置。自治体採用ではなく、事業者が業務委託として受託しても問題ありません。")]
fx=Inches(0.9); fw=Inches(5.7); fg=Inches(0.35); fy=Inches(2.3)
for i,(ti,amt,de) in enumerate(fin):
    x=Emu(int(fx)+i*(int(fw)+int(fg)))
    card=rect(s,x,fy,fw,Inches(3.0),RGBColor(0xF1,0xF6,0xFB),line=LINE,rounded=True)
    textbox(s,Emu(int(x)+int(Inches(0.4))),fy+Inches(0.35),Emu(int(fw)-int(Inches(0.8))),Inches(0.6),[(ti,18,NAVY,True)])
    textbox(s,Emu(int(x)+int(Inches(0.4))),fy+Inches(1.05),Inches(5),Inches(0.7),[(amt,32,ACCENT,True)])
    textbox(s,Emu(int(x)+int(Inches(0.4))),fy+Inches(1.9),Emu(int(fw)-int(Inches(0.8))),Inches(0.9),[(de,14,INK,False)])
    bdg=rect(s,Emu(int(x)+int(Inches(0.4))),fy+Inches(2.45),Inches(2.3),Inches(0.45),RGBColor(0xE6,0xF4,0xEC),rounded=True)
    shape_text(bdg,[[("✓ 総務省に確認済み",13,GREEN,True)]])
textbox(s,Inches(0),Inches(5.7),SW,Inches(0.5),
        [[("自治体様ごとに課題が異なるため、取り組む内容や予算に応じて柔軟にご相談可能です。",15,SUB,True)]],align=PP_ALIGN.CENTER)
footer(s,21)

# ============ 22 代表 ============
s=slide(); bg(s,WHITE)
kicker(s,"代表メッセージ")
photo=rect(s,Inches(0.9),Inches(1.5),Inches(3.4),Inches(5.0),NAVY2,rounded=True)
av=oval(s,Inches(1.9),Inches(2.1),Inches(1.4),Inches(1.4),RGBColor(0x35,0x5A,0x7A)); shape_text(av,[("◠",40,RGBColor(0xCF,0xE0,0xF0),True)])
textbox(s,Inches(1.1),Inches(5.0),Inches(3.0),Inches(0.6),[("佐近 航",24,WHITE,True)])
textbox(s,Inches(1.1),Inches(5.6),Inches(3.0),Inches(0.8),
        [[("株式会社KITAICHI 代表取締役",12,RGBColor(0xA9,0xC0,0xD6),False)],[("（さこん わたる）",12,RGBColor(0xA9,0xC0,0xD6),False)]])
bio=[("財政破綻を経験した夕張市役所で、地域再生を担当していました。そこで痛感したのが「地域のプレイヤー不足」です。良いアイデアや知恵があっても、それを実行する担い手が足りていません。",16,INK,False),
     ("",8,INK,False),
     ("その課題を解決するためにKITAICHIを創業。地域に新しいことを起こす人たちのクラウドファンディングを通じて、情報発信や応援者との継続的な関係づくりに取り組んでいます。",16,INK,False)]
textbox(s,Inches(4.7),Inches(1.7),Inches(8.0),Inches(2.8),[bio[0],bio[1],bio[2]])
hl=rect(s,Inches(4.7),Inches(4.7),Inches(8.0),Inches(1.5),LIGHT,rounded=True)
rect(s,Inches(4.7),Inches(4.7),Inches(0.08),Inches(1.5),ACCENT)
textbox(s,Inches(5.0),Inches(4.95),Inches(7.5),Inches(1.1),
        [("ふるさと住民登録制度は、地域の価値に共感する応援者の「質と量」を高め、プレイヤー不足を解決する制度。最大限活用するため、協働させてください。",16,NAVY,True)])
footer(s,22)

# ============ 23 スタッフ ============
s=slide(); bg(s,WHITE)
kicker(s,"運営スタッフ")
textbox(s,LX,Inches(1.1),Inches(11),Inches(0.7),[("私たちがサポートします。",30,NAVY,True)])
staff=[("佐近 航","行政経験とふるさと納税サイト運営の経験。地域創生はお任せください。"),
       ("宮岸 由海","ふるさと納税サイト運営を経験。学生のプロジェクトに多く関わっています。"),
       ("佐藤 千咲","准認定ファンドレイザー。地域イベントの企画運営にも従事しています。"),
       ("松代 空","宿・カフェ運営を経験。宿や飲食の立ち上げはぜひご相談を。")]
sx=Inches(0.85); sw_=Inches(2.85); sg=Inches(0.23); sy=Inches(2.5)
for i,(nm,de) in enumerate(staff):
    x=Emu(int(sx)+i*(int(sw_)+int(sg)))
    card=rect(s,x,sy,sw_,Inches(3.2),RGBColor(0xF6,0xF9,0xFC),line=LINE,rounded=True)
    av=oval(s,Emu(int(x)+int(sw_)/2-int(Inches(0.55))),sy+Inches(0.4),Inches(1.1),Inches(1.1),RGBColor(0xE7,0xEE,0xF5))
    shape_text(av,[("◠",30,NAVY2,True)])
    textbox(s,x,sy+Inches(1.65),sw_,Inches(0.4),[(nm,18,NAVY,True)],align=PP_ALIGN.CENTER)
    textbox(s,Emu(int(x)+int(Inches(0.25))),sy+Inches(2.15),Emu(int(sw_)-int(Inches(0.5))),Inches(1.0),
            [(de,13,SUB,False)],align=PP_ALIGN.CENTER)
footer(s,23)

# ============ 24 最後に (dark) ============
s=slide(); bg(s,NAVY)
oval(s,Inches(9.5),Inches(4.5),Inches(5),Inches(5),NAVY2)
textbox(s,Inches(0),Inches(1.9),SW,Inches(0.4),[("最後に",14,ACCENT2,True)],align=PP_ALIGN.CENTER)
textbox(s,Inches(1.0),Inches(2.5),Inches(11.3),Inches(2.0),
        [[("地域の担い手と応援者をつなぎ、",32,WHITE,True)],
         [("地域の活性化と担い手不足の解消に",32,WHITE,True)],
         [("貢献させてください。",32,WHITE,True)]],align=PP_ALIGN.CENTER)
textbox(s,Inches(2.4),Inches(5.2),Inches(8.5),Inches(1.2),
        [("KITAICHIは、クラウドファンディングを通じて地域の挑戦を全国へ届けてきました。その力を、ふるさと住民登録制度の運用に活かします。",16,RGBColor(0xCD,0xDD,0xF0),False)],align=PP_ALIGN.CENTER)
footer(s,24,dark=True)

# ============ 25 CTA (dark) ============
s=slide(); bg(s,NAVY)
textbox(s,Inches(0),Inches(1.5),SW,Inches(0.4),[("ご相談・お問い合わせ",14,ACCENT2,True)],align=PP_ALIGN.CENTER)
textbox(s,Inches(0),Inches(2.05),SW,Inches(0.7),[("まずは雑談レベルでも大丈夫です。",32,WHITE,True)],align=PP_ALIGN.CENTER)
pill=rect(s,Inches(3.67),Inches(3.2),Inches(6.0),Inches(0.95),ACCENT,rounded=True)
try: pill.adjustments[0]=0.5
except Exception: pass
shape_text(pill,[("まず30分のオンライン相談から",22,WHITE,True)])
textbox(s,Inches(2.4),Inches(4.4),Inches(8.5),Inches(1.0),
        [[("御地に合った活用アイデアを、無料でご提案します。",17,RGBColor(0xCD,0xDD,0xF0),False)],
         [("「うちの町だと何ができる？」を、一緒に考えるところから。",17,RGBColor(0xCD,0xDD,0xF0),False)]],align=PP_ALIGN.CENTER)
mail=rect(s,Inches(4.17),Inches(5.7),Inches(5.0),Inches(0.7),NAVY2,rounded=True)
shape_text(mail,[[("担当：佐近　　✉ info@kita1.jp",20,WHITE,True)]])
footer(s,25,dark=True)

prs.save("/home/user/kitaichi/furusato_proposal.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
