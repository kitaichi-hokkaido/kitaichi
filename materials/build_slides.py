# -*- coding: utf-8 -*-
"""Build a Google-Slides-ready PPTX that mirrors the A4 one-pager."""
import sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

OUT = sys.argv[1] if len(sys.argv) > 1 else 'KITAICHI_スライド.pptx'
IMG_STU = sys.argv[2] if len(sys.argv) > 2 else 'slide_students.jpg'
IMG_CF = sys.argv[3] if len(sys.argv) > 3 else 'slide_cf.jpg'
FONT = 'Noto Sans JP'

# palette
TEAL = (0x0a, 0x6e, 0x6e)
TEALD = (0x0a, 0x3d, 0x3d)
LIGHT = (0xf1, 0xf8, 0xf8)
GOLD = (0xff, 0xd8, 0x4d)
GOLD2 = (0xff, 0xe0, 0x8a)
INK = (0x1f, 0x29, 0x33)
GRAY = (0x5a, 0x6b, 0x6b)
AMBER_BG = (0xff, 0xf7, 0xec)
AMBER_LN = (0xf0, 0xd9, 0xb0)
AMBER_TX = (0xb8, 0x74, 0x1a)
WHITE = (0xff, 0xff, 0xff)
LINEGRY = (0xcb, 0xd6, 0xd6)

prs = Presentation()
prs.slide_width = Cm(21)
prs.slide_height = Cm(29.7)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes

AL = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
AN = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def rect(x, y, w, h, fill=None, line=None, lw=0.75, rounded=False):
    shp = S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                      Cm(x), Cm(y), Cm(w), Cm(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(*fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = RGBColor(*line); shp.line.width = Pt(lw)
    if rounded:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    return shp


def text(x, y, w, h, runs, size=9, color=INK, bold=False, align='l', anchor='t',
         ls=1.08, sp_after=0.0):
    """runs: str OR list of paragraphs; each paragraph is str OR list of (txt,{opts})."""
    tb = S.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = AN[anchor]
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Cm(0.08))
    tf.margin_top = Cm(0.02); tf.margin_bottom = Cm(0.02)
    if isinstance(runs, str):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = AL[align]
        p.line_spacing = ls
        if sp_after:
            p.space_after = Pt(sp_after)
        p.space_before = Pt(0)
        chunks = [para] if isinstance(para, str) else para
        for ch in chunks:
            if isinstance(ch, str):
                t, o = ch, {}
            else:
                t, o = ch
            r = p.add_run(); r.text = t
            f = r.font
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.name = FONT
            f.color.rgb = RGBColor(*o.get('color', color))
    return tb


def picture_fit(path, x, y, w, h):
    """Contain: fit whole image inside box, centered."""
    iw, ih = Image.open(path).size
    box_r = w / h
    img_r = iw / ih
    if img_r > box_r:
        nw = w; nh = w / img_r
    else:
        nh = h; nw = h * img_r
    px = x + (w - nw) / 2
    py = y + (h - nh) / 2
    S.add_picture(path, Cm(px), Cm(py), Cm(nw), Cm(nh))


M = 1.0           # side margin
W = 21 - 2 * M    # content width = 19

# ---------- Header ----------
text(M, 0.55, 12, 1.0, [[('株式会社KITAICHI ', {'size': 20, 'bold': True, 'color': TEALD}),
                         ('（キタイチ）', {'size': 10, 'color': GRAY})]], anchor='b')
text(M + 7, 0.5, W - 7, 1.0,
     [[('チャレンジする高校生は、', {'size': 10.5, 'bold': True, 'color': TEAL})],
      [('地域が変わるきっかけになる。', {'size': 10.5, 'bold': True, 'color': TEAL})]],
     align='r', anchor='b', ls=1.15)
rect(M, 1.65, W, 0.03, fill=TEAL)

# ---------- Photos ----------
py0 = 1.95
ph = 4.7
lw_ = 7.0
gap = 0.5
picture_fit(IMG_STU, M, py0, lw_, ph)
# CF image fills its cell (ratio matches)
rx = M + lw_ + gap
rw = W - lw_ - gap
picture_fit(IMG_CF, rx, py0, rw, ph)

# ---------- Title bar ----------
ty = py0 + ph + 0.25
rect(M, ty, W, 0.85, fill=TEAL, rounded=False)
text(M, ty, W, 0.85, '会社概要 ＆ クラウドファンディング活用サポートのご案内',
     size=13, bold=True, color=WHITE, align='c', anchor='m')

# ---------- Service heading ----------
hy = ty + 1.15
rect(M, hy + 0.02, 0.14, 0.5, fill=TEAL)
text(M + 0.3, hy, 10, 0.55, 'サービス概要', size=12.5, bold=True, color=TEALD, anchor='m')

# service lead bar
ly = hy + 0.68
rect(M, ly, W, 1.0, fill=LIGHT, rounded=True)
text(M + 0.3, ly, W - 0.6, 1.0,
     [[('高校生のチャレンジを、', {'size': 14, 'bold': True, 'color': TEALD}),
       ('クラウドファンディング', {'size': 14, 'bold': True, 'color': TEAL}),
       ('で実現します。', {'size': 14, 'bold': True, 'color': TEALD})]],
     anchor='m')

# intro
iy = ly + 1.15
text(M, iy, W, 1.2,
     [[('高校生が主体的に地域のために挑戦し、町内外から応援を集めて実現する経験は、', {'size': 9.5}),
       ('生徒の成長にも、地域の未来にも大きな意義があります。', {'size': 9.5, 'color': TEAL, 'bold': True}),
       ('「やりたいけれど予算がない」というハードルを、企画から支援獲得までの伴走で乗り越えます。提供内容は次の2つです。', {'size': 9.5})]],
     ls=1.25)

# ---------- Offer cards ----------
oy = iy + 1.6
ch = 2.0
cw = (W - 0.5) / 2
for idx, (cx, num, title, body) in enumerate([
    (M, '提供①', 'クラウドファンディングで地域での探究活動を実現',
     '生徒の「地域のためにやりたいこと」を、企画づくりから支援の獲得まで伴走し、実現まで導きます。（例：花火大会・地域イベント 等）'),
    (M + cw + 0.5, '提供②', '総合的な探究学習の伴走／単発授業にも対応',
     '年間を通じた探究の伴走のほか、1回〜数回の授業でクラウドファンディングを解説することも可能。予算に応じ柔軟に対応します。')]):
    rect(cx, oy, cw, ch, fill=LIGHT, line=(0xb9, 0xd4, 0xd4), lw=0.75, rounded=True)
    rect(cx, oy, cw, 0.09, fill=TEAL)
    rect(cx + 0.25, oy + 0.22, 1.45, 0.42, fill=TEAL, rounded=True)
    text(cx + 0.25, oy + 0.2, 1.45, 0.45, num, size=9, bold=True, color=WHITE, align='c', anchor='m')
    text(cx + 0.25, oy + 0.7, cw - 0.5, 0.6, title, size=10, bold=True, color=TEALD, ls=1.1)
    text(cx + 0.25, oy + 1.28, cw - 0.5, 0.7, body, size=8.5, color=(0x3a, 0x4a, 0x4a), ls=1.18)

# ---------- Strengths + Flow ----------
sy = oy + ch + 0.35
colw = (W - 0.5) / 2
# strengths
rect(M, sy + 0.02, 0.14, 0.45, fill=TEAL)
text(M + 0.3, sy, colw, 0.5, 'サポートの強み', size=11, bold=True, color=TEALD, anchor='m')
text(M, sy + 0.6, colw, 3.0,
     [[('① プロによる伴走', {'size': 9, 'bold': True, 'color': TEAL})],
      [('元自治体職員・元サイト運営者・ファンドレイザーが、ページ制作・返礼品設計・プロモーションを支援。', {'size': 8.3, 'color': (0x3a, 0x4a, 0x4a)})],
      [('② 進行管理で負担を軽減', {'size': 9, 'bold': True, 'color': TEAL})],
      [('企画から終了までスケジュールを管理し、できる限り負担のかからない運営を実現。', {'size': 8.3, 'color': (0x3a, 0x4a, 0x4a)})],
      [('③ 850件超の実績に基づく伴走力', {'size': 9, 'bold': True, 'color': TEAL})],
      [('密に連絡を取り、生徒のモチベーション面まで寄り添いサポートします。', {'size': 8.3, 'color': (0x3a, 0x4a, 0x4a)})]],
     ls=1.2, sp_after=2)
# flow
fx = M + colw + 0.5
rect(fx, sy + 0.02, 0.14, 0.45, fill=TEAL)
text(fx + 0.3, sy, colw, 0.5, 'サポートの流れ（5ステップ）', size=11, bold=True, color=TEALD, anchor='m')
flow = [('01 企画設計', 'ヒアリング・スケジュール作成・企画内容の整理'),
        ('02 ページ制作', '構成作成・文章入力・写真準備・トップ画像制作・審査対応'),
        ('03 リターン設計', '返礼品のアイデア出し・内容検討・ページへの反映'),
        ('04 プロモ準備', '告知方法のヒアリング・戦略策定・告知先リスト・例文提供'),
        ('05 公開後対応', '告知活動・改善アドバイス・支援者対応・リターン発送')]
paras = []
for a, b in flow:
    paras.append([(a + '　', {'size': 8.3, 'bold': True, 'color': TEALD}),
                  (b, {'size': 8.3, 'color': (0x3a, 0x4a, 0x4a)})])
text(fx, sy + 0.6, colw, 3.0, paras, ls=1.15, sp_after=3)

# ---------- Price ----------
pry = sy + 3.7
rect(M, pry, W, 1.0, fill=TEAL, rounded=True)
text(M + 0.35, pry, 3.2, 1.0,
     [[('サポート料金', {'size': 10, 'bold': True, 'color': WHITE})],
      [('成功報酬型', {'size': 7, 'color': (0xbf, 0xe3, 0xe3)})]], anchor='m', ls=1.0)
text(M + 3.7, pry, W - 4.0, 1.0,
     [[('集まった支援額 × ', {'size': 11, 'color': WHITE}),
       ('5％ ＋ 税', {'size': 13, 'bold': True, 'color': GOLD2}),
       ('　― 持ち出しの費用は一切ございません（成果が出た分のみのお支払い）', {'size': 8, 'color': (0xd6, 0xef, 0xef)})]],
     anchor='m')

# ---------- Track record ----------
tky = pry + 1.25
rect(M, tky, W, 2.15, fill=AMBER_BG, line=AMBER_LN, lw=0.75, rounded=True)
text(M + 0.35, tky + 0.18, W - 0.7, 1.9,
     [[('これまでの実績・取り組み', {'size': 9, 'bold': True, 'color': AMBER_TX})],
      [('● ふるさとチョイス クラウドファンディングのサイト運営業務を受託し、全国自治体 850プロジェクトの立ち上げをサポート（令和7年度）。', {'size': 8.4})],
      [('● 三重県立相可高校の総合的な探究学習を1年間伴走。生徒が地域貢献の取り組みを具体化し、クラファンで実現するまで支援。', {'size': 8.4})],
      [('● 北海道立上ノ国高校の生徒会が、町民も楽しめる学校祭を目指し花火大会を企画。募金とクラファンで幅広く応援を獲得。', {'size': 8.4})]],
     ls=1.3, sp_after=2)

# ---------- Company + Philosophy ----------
cy = tky + 2.45
# company table
rect(M, cy + 0.02, 0.14, 0.45, fill=TEAL)
text(M + 0.3, cy, colw, 0.5, '会社概要', size=11, bold=True, color=TEALD, anchor='m')
rows = [('会社名', '株式会社KITAICHI（キタイチ）'),
        ('代表者', '代表　佐近　航（さこん わたる）'),
        ('事業内容', 'まちづくり・地域創生に関する事業のサポート／クラウドファンディングの企画・運営伴走'),
        ('主な強み', 'CAMPFIREパートナー企業。元自治体職員・元サイト運営者・ファンドレイザーが在籍'),
        ('お問い合わせ', 'メール：info@kita1.jp')]
tb = S.add_table(len(rows), 2, Cm(M), Cm(cy + 0.6), Cm(colw), Cm(3.0)).table
tb.columns[0].width = Cm(2.7)
tb.columns[1].width = Cm(colw - 2.7)
tb.first_row = False
for ri, (k, v) in enumerate(rows):
    for ci, val in enumerate((k, v)):
        cell = tb.cell(ri, ci)
        cell.margin_left = Cm(0.12); cell.margin_right = Cm(0.1)
        cell.margin_top = Cm(0.03); cell.margin_bottom = Cm(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xee, 0xf5, 0xf5) if ci == 0 else RGBColor(*WHITE)
        p = cell.text_frame.paragraphs[0]
        p.line_spacing = 1.05
        r = p.add_run(); r.text = val
        r.font.size = Pt(7.6); r.font.name = FONT
        r.font.bold = ci == 0
        r.font.color.rgb = RGBColor(*TEALD) if ci == 0 else RGBColor(*INK)

# philosophy
phx = M + colw + 0.5
rect(phx, cy + 0.02, 0.14, 0.45, fill=TEAL)
text(phx + 0.3, cy, colw, 0.5, '私たちの想い ─ 高校生のチャレンジを応援したい',
     size=10, bold=True, color=TEALD, anchor='m')
rect(phx, cy + 0.6, colw, 3.0, fill=(0xf4, 0xf9, 0xf9), line=(0xcb, 0xe0, 0xe0), lw=0.75, rounded=True)
text(phx + 0.25, cy + 0.75, colw - 0.5, 2.8,
     [[('高校生には、地域を変えるほどの力がある。', {'size': 9.5, 'bold': True, 'color': TEAL})],
      [('私は財政破綻を経験した夕張市役所で、地域再生や高校魅力化事業を担当していました。夕張高校の生徒会長が「夕張に言い訳するのではなく、夕張だからできることに目を向けたい」とまちづくり会議で大人たちに向けて話し、その言葉が財政再生計画の見直しにつながりました。高校生は地域を変える力があります。だからこそ、高校生のチャレンジに資金と応援を届けるサポートをさせてください。', {'size': 8.2})],
      [('株式会社KITAICHI 代表　佐近　航', {'size': 7.5, 'color': GRAY})]],
     ls=1.28, sp_after=2)

# ---------- Contact ----------
coy = cy + 3.95
rect(M, coy, W, 1.05, fill=TEALD, rounded=True)
text(M + 0.4, coy, 11, 1.05,
     [[('ご相談・お問い合わせ', {'size': 10, 'bold': True, 'color': WHITE})],
      [('まずは雑談レベルでもOKです。お気軽にご相談ください。', {'size': 7.5, 'color': (0xbb, 0xcc, 0xdd)})]],
     anchor='m', ls=1.05)
text(M + 11, coy, W - 11.4, 1.05,
     [[('株式会社KITAICHI', {'size': 9, 'color': WHITE})],
      [('メール：', {'size': 9, 'color': WHITE}), ('info@kita1.jp', {'size': 10, 'bold': True, 'color': WHITE})]],
     align='r', anchor='m', ls=1.1)

# ---------- Notes ----------
ny = coy + 1.15
text(M, ny, W, 1.2,
     [[('※上記成功報酬とは別に、クラウドファンディングサイト（CAMPFIRE等）の手数料がかかります。契約先は学校・保護者会などの任意団体でも対応可能です。', {'size': 7})],
      [('※基本はオンライン対応。交通費のご負担で対面授業にも参加いたします。授業の運営委託をご希望の場合は予算に合わせて柔軟に対応しますのでご相談ください。', {'size': 7})]],
     color=GRAY, ls=1.25)

prs.save(OUT)
print('saved', OUT)
